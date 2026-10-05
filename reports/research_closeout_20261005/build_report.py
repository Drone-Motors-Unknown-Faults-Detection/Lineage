"""將封存指標編成文件；不載入特徵、模型，也不計算新研究成績。"""
import argparse
import gzip
import hashlib
import json
import re
import subprocess
import zipfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
M_PROTOCOL = ROOT / 'output/fault_type_self_challenging_lock/2026-10-04-22-01-49/protocol.json'
M_PATHS = {
    'protocol': M_PROTOCOL,
    'fit': ROOT/'output/fault_type_self_challenging_fit/2026-10-04-22-03-00/locked_study.json',
    'source_proof': ROOT/'output/fault_type_self_challenging_source_verify/2026-10-04-22-09-26/source_verified.json',
    'evaluate': ROOT/'output/fault_type_self_challenging_evaluate/2026-10-04-22-16-34/evaluation.json.gz',
    'verify': ROOT/'output/fault_type_self_challenging_verify/2026-10-04-22-18-59/verified.json.gz',
    'report': ROOT/'output/fault_type_self_challenging_report/2026-10-05-22-17-39/summary.json.gz',
    'backup': ROOT/'output/fault_type_fixed_delivery/2026-10-05-22-32-32/member_verification.json',
}
BATCHES = {
 'literature_expansion':('文獻 C','C', 'fault_type_literature_study','30 分類配置與分數支線'),
 'mechanism_research_v2':('機制 D/P','D','fault_type_mechanism_study','解耦、RMD 係數與分塊'),
 'continuous_research':('Q','Q','fault_type_continuous_study','gain、形狀與對角度量'),
 'metric_classification_v1':('E','E','fault_type_metric_classification','度量分類、鄰居 energy、一／三中心'),
 'local_fisher_v1':('F','F','fault_type_local_fisher','PCA／LFDA10／20'),
 'discriminative_prototypes_v1':('G','G','fault_type_discriminative_prototypes','固定／GLVQ／anchor，一／三中心'),
 'smooth_l1_v1':('H','H','fault_type_smooth_l1','smooth-L1 Q／S 幾何'),
 'context_prototypes_v1':('I','I','fault_type_context_prototypes','RPM degree1／2，static／GLVQ／aux'),
 'context_rejection_v1':('J','J','fault_type_context_rejection','distance／ambiguity／OR 拒絕'),
 'axis_invariant_kernel_v1':('K','K','fault_type_axis_kernel','S3 核 alpha=0/.5/1'),
 'rpm_risk_extrapolation_v1':('L','L','fault_type_risk_extrapolation','REx beta=0/1/10'),
 'feature_self_challenging_v1':('M','M','fault_type_self_challenging','兩表示與四遮蔽控制'),
}


def read(path):
    path = Path(path)
    with (gzip.open(path, 'rt', encoding='utf-8') if path.suffix == '.gz' else path.open(encoding='utf-8')) as f:
        return json.load(f)


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024*1024), b''): h.update(block)
    return h.hexdigest()


def dump(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def pct(x):
    return 'NA' if x is None else f'{100*x:.2f}'


def decimal(x):
    return 'NA' if x is None else f'{x:.4f}'


def table(headers, rows):
    def cell(v): return str(v).replace('|', '／').replace('\n', ' ')
    return '\n'.join(['|'+ '|'.join(headers)+'|', '|'+ '|'.join(['---']*len(headers))+'|',
                       *['|'+'|'.join(cell(v) for v in row)+'|' for row in rows]])


def macro(method):
    if 'motor_macro_seed0' in method: return method['motor_macro_seed0']
    # 早期定義不將 healthy unknown alarm 冒充完整健康誤報。
    source = method.get('equal_motor_descriptive', {})
    keys = {'known_accuracy':'known_classification.accuracy',
            'fault_accuracy':'known_fault_classification.accuracy',
            'fault_macro_f1':'known_fault_classification.macro_f1',
            'unknown_recall':'unknown_rejection.unknown_recall'}
    return {key:source.get(old, {}).get('mean') for key, old in keys.items()}


def metric_row(key, method):
    x = macro(method)
    return [key,pct(x.get('known_accuracy')),pct(x.get('fault_accuracy')),
            decimal(x.get('fault_macro_f1')),pct(x.get('healthy_total_alarm')),pct(x.get('unknown_recall'))]


def build_content(inventory, output):
    output.mkdir(parents=True, exist_ok=True)
    catalog = read(inventory/'method_catalog.json.gz')
    checkpoints = read(inventory/'checkpoint_inventory.json')
    decision = read(inventory/'decision_record.json')
    supports = read(inventory/'fold_support.json')
    refs = read(HERE/'references.json')
    for group,item in catalog.items():
        if sha(item['summary']) != item['summary_sha256']:
            raise ValueError('原摘要 SHA 改變：'+group)
        if not (ROOT/'experiments'/(BATCHES[group][2]+'.py')).is_file():
            raise ValueError('程式入口不存在：'+group)
    m = catalog['feature_self_challenging_v1']
    original_m = read(m['summary'])
    if sha(m['summary']) != m['summary_sha256']: raise ValueError('M 摘要 SHA 改變')
    if decision['path'] != 'B' or any(x['main_screen'] != 'FAILED' for x in m['reliability'].values()):
        raise ValueError('本收尾稿僅適用全方法未通過的 B 判定')
    if original_m['completed_runs'] != 216 or len(original_m['paired_differences']) != 1134:
        raise ValueError('M 評估／配對數不符')
    substitutions = {}
    folds = []
    class_rows = []
    motors = ['T3','T1','T2']
    for item, train, cal, test in zip(supports,['T1','T2','T3'],['T2','T3','T1'],motors):
        roles = item['roles']
        unknown = sum(roles['test']['classes'][k] for k in m['protocol']['unknown_labels'])
        folds.append([item['fold'],train+'/'+str(roles['train']['samples']),cal+'/'+str(roles['calibration']['samples']),
                      test+'/'+str(roles['test']['samples']), roles['test']['samples']-unknown,unknown])
        for label, count in sorted(roles['test']['classes'].items()): class_rows.append([test,label,count])
    if sum(x['roles']['test']['samples'] for x in supports) != 28910: raise ValueError('支持筆數不符')
    substitutions['fold_table'] = table(['fold','train motor／筆','cal motor／筆','test motor／筆','test known','test unknown'],folds)
    substitutions['support_table'] = table(['test motor','配置','支持筆數'],class_rows)
    substitutions['control_table'] = table(['方法','known %','fault %','條件 F1','健康總誤報 %','unknown recall %'],
                                          [metric_row(k,m['methods'][k]) for k in ['C02','C17','C24','D01']])
    reject_rows = []
    for key in ['C02','C17','C24','D01']:
        x = macro(m['methods'][key]); final = m['methods'][key]['final_fault_all_samples']
        reject_rows.append([key,decimal(final['macro_f1']),decimal(x['unknown_auroc']),decimal(x['unknown_aupr']),
                            pct(x['coverage']),pct(x['known_correct_after_rejection'])])
    substitutions['control_rejection_table'] = table(['方法','全分母 final fault F1','unknown AUROC','unknown AUPR','coverage %','known拒絕後正確 %'],reject_rows)
    d,c,c17 = (macro(m['methods'][k]) for k in ['D01','C02','C17'])
    substitutions['control_delta'] = (f"D01 相對 C02：known accuracy +{100*(d['known_accuracy']-c['known_accuracy']):.2f} 個百分點，"
       f"fault accuracy +{100*(d['fault_accuracy']-c['fault_accuracy']):.2f} 個百分點，健康總誤報下降 {100*(c['healthy_total_alarm']-d['healthy_total_alarm']):.2f} 個百分點；"
       f"unknown recall 不變。相對 C17，unknown recall 提高 {100*(d['unknown_recall']-c17['unknown_recall']):.2f} 個百分點，但 known rejection 從 {pct(c17['known_rejection_rate'])}% 增至 {pct(d['known_rejection_rate'])}%，"
       f"known 拒絕後正確率從 {pct(c17['known_correct_after_rejection'])}% 降至 {pct(d['known_correct_after_rejection'])}%。D01 仍有零召回與高誤報工況，不是可靠 winner。")
    batch_rows = []
    for group,x in catalog.items():
        counts = x['counts']; completed=counts.get('completed_runs',counts.get('completed_evaluations'))
        planned=counts.get('planned_evaluations', len(x['protocol']['arms'])*9)
        gates=Counter(t['main_screen'] for t in x['reliability'].values())
        status = '／'.join(f'{k}:{v}' for k,v in gates.items()) or '無同版主篩；保留探索結果'
        batch_rows.append([BATCHES[group][0],planned,completed,planned-completed,counts.get('prediction_records'),status])
    substitutions['batch_table'] = table(['批次','planned評估','completed','INCOMPLETE','預測records','可靠性'],batch_rows)
    selected=[('continuous_research','Q07'),('local_fisher_v1','F02'),('discriminative_prototypes_v1','G16'),
              ('context_rejection_v1','J03'),('axis_invariant_kernel_v1','K09'),('rpm_risk_extrapolation_v1','L10'),
              ('rpm_risk_extrapolation_v1','L16'),('feature_self_challenging_v1','M13'),('feature_self_challenging_v1','M22')]
    substitutions['selected_table']=table(['方法','known %','fault %','條件 F1','健康總誤報 %','unknown recall %'],
                                          [metric_row(key,catalog[group]['methods'][key]) for group,key in selected])
    seed_rows=[]
    for key in [arm['id'] for arm in m['protocol']['arms']]:
        for seed,values in m['methods'][key]['per_seed_motor_macro'].items():
            seed_rows.append([key+'/'+seed,pct(values['fault_accuracy']),decimal(values['fault_macro_f1']),
                              pct(values['healthy_total_alarm']),pct(values['unknown_recall'])])
    substitutions['m_seed_table']=table(['M方法／seed','fault %','條件 F1','健康總誤報 %','unknown recall %'],seed_rows)
    mm=macro(m['methods']['M13'])
    substitutions['m_delta']=(f"M13 相對 C17 的 fault accuracy 提高 {100*(mm['fault_accuracy']-c17['fault_accuracy']):.2f} 個百分點，"
        f"條件 F1 變動 {mm['fault_macro_f1']-c17['fault_macro_f1']:+.4f}，健康總誤報增加 {100*(mm['healthy_total_alarm']-c17['healthy_total_alarm']):.2f} 個百分點。"
        f"這個局部分類增加不符合 F1 和安全保護。M13 相對 C24 的 fault accuracy 只增加 {100*(mm['fault_accuracy']-macro(m['methods']['C24'])['fault_accuracy']):.2f} 個百分點，未達 +2 個百分點。")
    motor_rows=[];rpm_rows=[];recalls=[]
    for x in m['methods']['M13']['per_motor']:
        v=x['metrics'];motor_rows.append([x['motor'],pct(v['known_accuracy']),pct(v['fault_accuracy']),decimal(v['fault_macro_f1']),pct(v['healthy_total_alarm']),pct(v['unknown_recall'])])
        for cell in x['rpms']:
            z=cell['metrics'];health=z['healthy_safety_v2'];u=z['unknown_rejection']
            rpm_rows.append([x['motor']+'/'+cell['rpm'],health['n_healthy'],pct(health['healthy_to_unknown_rate']),
                             pct(health['healthy_to_known_fault_rate']),pct(health['healthy_total_alarm_rate']),pct(u['unknown_recall'])])
        for cl in x['classes']: recalls.append([x['motor'],cl['label'],cl['samples'],pct(cl['classifier_recall']),pct(cl['rejection_rate'])])
    substitutions['m_motor_table']=table(['M13 test motor','known %','fault %','條件 F1','健康總誤報 %','unknown recall %'],motor_rows)
    substitutions['m_rpm_table']=table(['M13 motor／RPM','健康筆數','健康→未知 %','健康→已知fault %','健康總誤報 %','unknown recall %'],rpm_rows)
    substitutions['m_class_table']=table(['M13 motor','配置','支持筆數','classifier recall %','reject rate %'],recalls)
    gate_rows=[];reason_rows=[]
    for key, gate in m['reliability'].items():
        method=m['methods'][key]
        gate_rows.append([key,gate['main_screen'],pct(method['worst_motor_fault_class_recall']),
                          pct(method['worst_rpm_healthy_total_alarm']),pct(method['worst_motor_unknown_recall'])])
        reason_rows.append(f"### {key} 契約未通過原因\n\n"+'；'.join(gate['reasons'])+'。\n\nR2 多 subset／壓力證據未完成；fresh final test 與每類群組數 INCOMPLETE，採集來源 UNKNOWN。')
    substitutions['m_gate_table']=table(['M方法','主篩','最差類召回 %','最差工況健康誤報 %','最差motor未知召回 %'],gate_rows)
    substitutions['m_reasons']='\n\n'.join(reason_rows)
    substitutions['status_table']=table(['項目','狀態','支持範圍'],[
        ['正式計算鏈／逐筆預測','VERIFIED','來源 SHA／IDs、known-only fit、216逐筆驗證與1,134配對'],
        ['M可靠提升','FAILED','24方法未達原契約；沒有部署替換'],
        ['多known subsets穩健改善','INCOMPLETE','近期初篩只用一組N=5；歷史Nsweep不替代新方法驗收'],
        ['獨立fresh／每類≥2 test groups','INCOMPLETE','所有rows歷史曝露，每折一顆test motor'],
        ['raw sessions／overlap／清理遮罩','UNKNOWN','formal副本檢查不能證明採集獨立'],
        ['N Group DRO','未實作','只有手冊與來源卡；不是演算法FAILED'],
        ['Word逐頁排版','見content_qa.json','內容與ZIP檢查不等於PNG視覺驗收']])
    parts=[]
    arm_inventory=[]
    for group,x in catalog.items():
        name,_,module,desc=BATCHES[group]
        p=x['protocol'];arms={v['id']:v for v in p['arms']}
        parts.append(f"### {name} {desc}\n\n程式入口 experiments/{module}.py；參數及 sources 沿 {x['summary']} 與 catalog[{group}].protocol。訓練／校準／測試 motor 按正文三折固定。沒有 global winner。")
        configs=[];metrics=[]
        for key, method in x['methods'].items():
            arm_id=key.split('/')[0]; arm=arms.get(arm_id)
            if arm is None: continue
            configs.append([key,'；'.join(f'{k}={v}' for k,v in arm.items() if k!='id')])
            metrics.append(metric_row(key,method))
            arm_inventory.append(dict(batch=group,method_id=key,arm=arm,parameters=p.get('parameters',{}),
                                      source_documents=x['source_documents'],module=f'experiments/{module}.py',
                                      execution_status='INCOMPLETE' if macro(method).get('known_accuracy') is None and not method.get('per_motor') else '已保存實測',
                                      summary=x['summary'],summary_sha256=x['summary_sha256']))
        parts.append(table(['方法 ID','固定配置'],configs))
        parts.append(table(['方法','known %','fault %','條件 F1','完整健康誤報 %','unknown recall %'],metrics))
    substitutions['all_methods']='\n\n'.join(parts)
    dump(output/'arm_inventory.json',arm_inventory)
    substitutions['references']='\n\n'.join(f"[{v['id']}] {v['citation']} [原始來源]({v['url']})\n\n閱讀範圍：{v['depth']}。本站用途：{v['use']}。" for v in refs)
    artifact_rows=[];artifacts={}
    for key,path in M_PATHS.items():
        if not path.is_file(): raise ValueError('M 來源缺失：'+str(path))
        digest=sha(path);artifacts[key]=dict(path=str(path),sha256=digest,bytes=path.stat().st_size)
        artifact_rows.append([key,str(path.relative_to(ROOT)),digest])
    substitutions['artifact_table']=table(['M stage','實際索引路徑','實體 SHA256'],artifact_rows)
    substitutions['reproduction']=('研究根目錄：'+str(ROOT)+'。備份根目錄：D:/schoolshit/專題/src/lineage_fault_type_artifacts/2026-10-04/feature_self_challenging_v1/archives。\n\n'
       '唯讀盤點：`.venv310/Scripts/python.exe -m experiments.fault_type_research_closeout --m-report output/fault_type_self_challenging_report/2026-10-05-22-17-39/summary.json.gz`。\n\n'
       'M 既有 report 重算：`.venv310/Scripts/python.exe -m experiments.fault_type_self_challenging report --protocol output/fault_type_self_challenging_lock/2026-10-04-22-01-49/protocol.json --evaluation output/fault_type_self_challenging_evaluate/2026-10-04-22-16-34/evaluation.json.gz --verification output/fault_type_self_challenging_verify/2026-10-04-22-18-59/verified.json.gz --baseline output/fault_type_metrics_v2/2026-10-02-12-51-52/metrics_v2.json --previous output/fault_type_mechanism_evaluate/2026-10-02-17-53-06/evaluation.json`。\n\n'
       '備份核對：`.venv310/Scripts/python.exe -m experiments.fault_type_fixed_delivery --index output/fault_type_self_challenging_backup/2026-10-05-22-24-22/archive_index.json --index output/fault_type_archive/2026-10-05-22-28-54/archive_index.json`。僅核對，不解包覆寫。\n\n'
       '恢復時先依各批 artifact_location.json 尋找 primary_artifact；核對 ZIP whole SHA／member SHA／CRC，再在新的空目錄還原。沿各批 result_index 的 parent_protocol／parent_lock 復原父 reference；不得用 compact 摘要冒充完整模型或逐樣本 CSV。')
    substitutions['delivery_notes']=(f"盤點根目錄 {inventory}；當前內容產物 {output}。checkpoint_inventory、decision_record、fold_support、method_catalog.json.gz 與 arm_inventory.json 是可重算／追溯入口。"
       ' accuracy study 背景另見 reports/fault_type_accuracy_study/result_index.json：198新評估、18 A0重用，不與主研究controls重複加總；老師歷史2,490另列。\n\n'
       '本回合先行手冊 4fb72015d8461846d90f19dd9fac7e25dcd56a8e、盤點63f1e2e9a161b31e4a5a8c23c7d31a5918770630 均已push並核對remote。M協定5d4893b7ac89d4570cc764703952354841058592、fit ffbc0aac9bea74f70f5f4bd151653a3b0f7efa07、source e0798d72254b03bd3177aa5dd00d965cfec6a93f 均已保存；後續交付commit以execution_log與delivery_index為準，避免把文件自己的SHA循環寫入文件。')
    body=(HERE/'report_body.md').read_text(encoding='utf-8')
    for key,value in substitutions.items(): body=body.replace('{{'+key+'}}',value)
    if re.search(r'\{\{\w+\}\}',body): raise ValueError('文件占位符未展開')
    (output/'research_report.md').write_text(body,encoding='utf-8')
    m_findings = '# 有限表格自挑戰研究收尾\n\n2026-10-05，Asia/Taipei。B 路徑；216 completed、0 failed runs、1,134 配對、24 方法全部 FAILED。\n\n'
    m_findings += substitutions['m_seed_table']+'\n\n'+substitutions['m_gate_table']+'\n\n'+substitutions['m_reasons']
    (output/'m_final_findings.md').write_text(m_findings,encoding='utf-8')
    delivery=dict(batch='exp21_feature_self_challenging_v1',path='B',completed=216,failed_runs=0,
                  prediction_records=2081520,unique_samples=28910,paired_comparisons=1134,reliability_failed_methods=24,
                  artifacts=artifacts,semantic_seals={k:original_m[k] for k in
                      ['protocol_checksum','evaluation_checksum','verification_checksum','report_checksum']},
                  backup_indices=['output/fault_type_self_challenging_backup/2026-10-05-22-24-22/archive_index.json',
                                  'output/fault_type_archive/2026-10-05-22-28-54/archive_index.json'],
                  backup_verification=artifacts['backup'],decision_record=str(inventory/'decision_record.json'),
                  findings=str(output/'m_final_findings.md'),inventory=str(inventory),
                  group_dro='已規劃未實作，本階段收尾，未進入訓練或測試',
                  production_replacement=False,reliable_improvement_reached=False,fresh_final_test=False)
    dump(output/'m_result_index.json',delivery)
    # 只檢查本次新中文稿；保留來源 ID、URL 與英文 reasons。
    dedicated='补称没复为这却仅学类训练后据样数录验论术对开关体问该笔达当拟损归况简阅页实际从应门节认报败遗误厂测泄优标强经区华属时观选变权线'
    chinese_only=re.sub(r'https?://\S+|`[^`]*`','',body)
    unexpected=sorted(set(chinese_only)&set(dedicated))
    qa=dict(numbers_from_sealed_summaries=True,matching_m_rows=72,fold_test_support=28910,
            m_completed=216,paired=1134,gate_counts=dict(Counter(g['main_screen'] for g in m['reliability'].values())),
            simplified_candidates=unexpected,layout_status='NOT_RENDERED',historical_full_tests_rerun=False,
            source_catalog_sha256=sha(inventory/'method_catalog.json.gz'),markdown_sha256=sha(output/'research_report.md'))
    dump(output/'content_qa.json',qa)
    return body,qa


def make_word(markdown, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(11 if name=='Normal' else 22 if name=='Title' else 15 if name=='Heading 1' else 12)
        s.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
        s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
        s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.15
    normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Lineage 研究收尾　')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def text(p,value):
        # 原生 hyperlink 與可編輯文字，移除 Markdown 行內 backticks。
        value=value.replace('`','')
        cursor=0
        for match in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)',value):
            p.add_run(value[cursor:match.start()]);link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),p.part.relate_to(match[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=match[1];run.append(t);link.append(run);p._p.append(link);cursor=match.end()
        p.add_run(value[cursor:])
    def word_table(lines):
        headers=[x.strip() for x in lines[0].strip('|').split('|')]
        rows=[[x.strip() for x in line.strip('|').split('|')] for line in lines[2:]]
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        width=6.87/len(headers)
        for col in t.columns: col.width=Inches(width)
        for cell,v in zip(t.rows[0].cells,headers):text(cell.paragraphs[0],v)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,v in zip(t.add_row().cells,row):text(cell.paragraphs[0],v)
        for row in t.rows:
            for cell in row.cells:
                cell.width=Inches(width);pr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'70');x.set(qn('w:type'),'dxa');margins.append(x)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                    for r in p.runs:r.font.size=Pt(10.5)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','bottom','left','right','insideH','insideV']:
            x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
        t._tbl.tblPr.append(borders)
        doc.add_paragraph()
    lines=markdown.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            word_table(block);continue
        if not line.strip():i+=1;continue
        if line.startswith('EQUATION:'):
            p=doc.add_paragraph();math=OxmlElement('m:oMath');run=OxmlElement('m:r');t=OxmlElement('m:t')
            t.text=line.partition(':')[2].strip();run.append(t);math.append(run);p._p.append(math)
        elif line.startswith('# '):text(doc.add_paragraph(style='Title'),line[2:])
        elif line.startswith('## '):text(doc.add_paragraph(style='Heading 1'),line[3:])
        elif line.startswith('### '):text(doc.add_paragraph(style='Heading 2'),line[4:])
        else:text(doc.add_paragraph(),line)
        i+=1
    doc.core_properties.title='單顆馬達訓練與跨馬達校準的開放集故障辨識研究'
    doc.core_properties.author='Codex'
    doc.core_properties.subject='Lineage 三馬達固定用途研究收尾'
    path=output/'research_report.docx';doc.save(path)
    # 結構核對只證明可編輯內容，不冒充視覺驗收。
    with zipfile.ZipFile(path) as z:
        if z.testzip() is not None:raise ValueError('DOCX ZIP CRC 不符')
        xml=z.read('word/document.xml').decode('utf-8')
        relationships=z.read('word/_rels/document.xml.rels').decode('utf-8')
        if xml.count('<m:oMath>')!=5:raise ValueError('原生公式數不符')
        if xml.count('<w:tbl>')<20:raise ValueError('表格缺失')
        if len(doc.tables)!=xml.count('<w:tbl>'):raise ValueError('表格結構不符')
        for token in ['37.29','31.25','0.2666','30.46','18.44','FAILED','INCOMPLETE','UNKNOWN']:
            if token not in xml:raise ValueError('Word 關鍵值缺失：'+token)
        if 'word/media/' in '\n'.join(z.namelist()):raise ValueError('不應以圖片代替文字')
    return dict(docx=str(path),docx_sha256=sha(path),editable_tables=len(doc.tables),native_equations=5,
                hyperlink_count=relationships.count('TargetMode="External"'),zip_crc='PASS',
                structural_status='PASS',layout_status='NOT_RENDERED')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inventory',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--word',action='store_true')
    args=parser.parse_args();body,qa=build_content(args.inventory,args.output)
    if qa['simplified_candidates']:raise ValueError('新稿仍有待人工核對的簡體候選：'+str(qa['simplified_candidates']))
    if args.word:qa.update(make_word(body,args.output))
    dump(args.output/'content_qa.json',qa)
    print(json.dumps(dict(output=str(args.output),word=args.word,qa=qa),ensure_ascii=False))


if __name__=='__main__':main()
